import json

from openpyxl import Workbook

from npl_sim.league import load_league_file
from npl_sim.rosters.npl_app import RosterTier
from npl_sim.rosters.sources import players_from_dump
from npl_sim.rosters.workbook import players_from_workbook, team_for_tab, write_dump

ALIGNMENT = load_league_file().for_season(2026)
TEAMS = [t for lg in ALIGNMENT.leagues for d in lg.divisions for t in d.teams]


def _tab(wb, title, division, rows):
    ws = wb.create_sheet(title)
    ws.append([None, None, None, None, division])
    ws.append([None, "ACTIVE ROSTER"])
    ws.append([None, "STARTING PITCHERS", "SS#", "MLB#"])
    for row in rows:
        ws.append(row)
    ws.append([None, "FINANCIALS"])
    ws.append([1, "Not, Aplayer", 1, 999, "C", "NYY"])  # past the roster: ignored


def _workbook(path):
    wb = Workbook()
    wb.active.title = "Key"
    for team in TEAMS:
        nick = team.split()[-1]
        rows = [[1, "Arm, Some", 1, 100000 + len(wb.sheetnames), "RHP", "NYY"]]
        if team == "Midwest Ducksnorts":
            rows = [
                [1, "Cole, Gerrit", 156, 543037.0, "RHP", "NYY"],
                ["x", "Milner, Hoby", 1245, 571948.0, "LHP", "CHC"],
                [None, "INFIELDERS"],
                [1, "Keys, Sean", "-", 815873.0, "3B", "TOR"],
                [None, "7-DAY INJURED LIST", None, None, None, None, ""],
                [1, "Winn, Keaton", 1346, 676775.0, "RHP", "SF"],
                [None, "TRIPLE-A ALABAMA QUACKS"],
                [None, "ON OPTION"],
                [1, "Gervase, Paul", 1353, 801434.0, "RHP", "LAD"],
                [None, "ASSIGNED OUTRIGHT"],
                ["-", "Cortes, Nestor", "-", 641482.0, "LHP", "SD"],
                [None, "SINGLE-A LAKELAND WATERFOWL"],
                ["-", "Dohm, Nate", "-", 702296.0, "RHP", "STL"],
            ]
        _tab(
            wb, f"{nick} #{len(wb.sheetnames)}", "AMERICAN LEAGUE\nSPLENDID SPLINTER DIVISION", rows
        )
    wb.create_sheet("IDs")
    wb.save(path)


def test_tabs_match_config_teams():
    assert team_for_tab("Ducksnorts #12", TEAMS) == "Midwest Ducksnorts"
    assert team_for_tab("Guys #21", TEAMS) == "Bad Guys"
    assert team_for_tab("Rats #8", TEAMS) == "Absolute Sickos"
    assert team_for_tab("Nobody #99", TEAMS) is None


def test_workbook_sections_become_npl_app_flags(tmp_path):
    _workbook(tmp_path / "rosters.xlsx")
    book = players_from_workbook(tmp_path / "rosters.xlsx", ALIGNMENT)
    assert len(book.tabs) == 24
    ducks = {p.name: p for p in book.players if p.team == "Midwest Ducksnorts"}
    assert ducks["Gerrit Cole"].tier is RosterTier.MLB_ROSTER
    assert ducks["Gerrit Cole"].mlb_id == 543037 and ducks["Gerrit Cole"].scoresheet_id == "156"
    assert ducks["Sean Keys"].tier is RosterTier.MLB_ROSTER and ducks["Sean Keys"].position == "3B"
    assert ducks["Hoby Milner"].tier is None  # "x": listed active but off the 40-man
    assert ducks["Keaton Winn"].tier is RosterTier.IL_7
    assert ducks["Paul Gervase"].tier is RosterTier.AAA
    assert ducks["Nestor Cortes"].tier is RosterTier.OUTRIGHTED
    assert ducks["Nate Dohm"].tier is RosterTier.A
    assert "Aplayer Not" not in ducks
    assert book.divisions["Midwest Ducksnorts"] == "AL Splendid Splinter"
    assert not any(m.startswith("Midwest Ducksnorts:") for m in book.division_mismatches(ALIGNMENT))
    assert any(m.startswith("New York Valkyries:") for m in book.division_mismatches(ALIGNMENT))

    # The dump reads back exactly like the NPL app's own export.
    path = write_dump(book.players, tmp_path / "npl.json")
    back = {p.mlb_id: p for p in players_from_dump(path)}
    for p in book.players:
        assert back[p.mlb_id].tier == p.tier and back[p.mlb_id].team == p.team
    assert all(r["model"] in ("npl.team", "npl.player") for r in json.loads(path.read_text()))


def test_committed_league_rosters_cover_every_team():
    by_team = {}
    for p in players_from_dump("data/rosters/npl-2026.json"):
        if p.is_active_mlb:
            by_team.setdefault(p.team, []).append(p)
    assert sorted(by_team) == sorted(TEAMS)
    assert all(25 <= len(r) <= 40 for r in by_team.values())
    ids = {p.mlb_id for r in by_team.values() for p in r}
    assert 805347 in ids and 80537 not in ids  # Jim Jarvis, fixed (MLBAM_FIXES)
