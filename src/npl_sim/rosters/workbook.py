"""Interim roster source: the league's roster workbook (one tab per team).

Dave (2026-10-12): pour in the league's current rosters so full-season sims
can be tested before the NPL app's database login is issued. The NPL app
(jeremyjbowers/npl) loads this same workbook into its own tables, so this
reader produces the same `NplPlayer` records and the rest of the sim can't
tell the two sources apart. When the app's login arrives, `sources.py`
replaces this module and nothing else changes.

Each team tab lists players under section headers in column B (ACTIVE
ROSTER with its STARTING PITCHERS / RELIEF PITCHERS / INFIELDERS /
OUTFIELDERS groups, the injured lists, TRIPLE-A ... with ON OPTION and
ASSIGNED OUTRIGHT, FOREIGN, RETIRED, NON-ROSTER, DOUBLE-A ..., SINGLE-A ...).
A player row has the 40-man marker in column A (1 = on the 40-man), the
name "Last, First" in B, Scoresheet ID in C, MLBAM ID in D, position in E
and MLB organization in F. FINANCIALS ends the roster.
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from pathlib import Path

from npl_sim.league import Alignment
from npl_sim.rosters.npl_app import NplPlayer, parse_mlb_id

# Section header (prefix) -> NPL app roster flag. ACTIVE ROSTER groups carry none.
SECTIONS = (
    ("7-DAY INJURED LIST", "roster_7dayIL"),
    ("56-DAY INJURED LIST", "roster_56dayIL"),
    ("END OF SEASON INJURED LIST", "roster_eosIL"),
    ("RESTRICTED LIST", "roster_restricted"),
    ("TRIPLE-A", "roster_tripleA"),
    ("ON OPTION", "roster_tripleA_option"),
    ("ASSIGNED OUTRIGHT", "roster_outrighted"),
    ("FOREIGN", "roster_foreign"),
    ("RETIRED", "roster_retired"),
    ("NON-ROSTER", "roster_nonroster"),
    ("DOUBLE-A", "roster_doubleA"),
    ("SINGLE-A", "roster_singleA"),
)
ACTIVE_HEADERS = (
    "ACTIVE ROSTER",
    "STARTING PITCHERS",
    "RELIEF PITCHERS",
    "INFIELDERS",
    "OUTFIELDERS",
)
END_HEADER = "FINANCIALS"

# MLBAM IDs the workbook has wrong, by the ID it lists. Jim Jarvis (Bears,
# SS, ATL) is listed as 80537, which the MLB Stats API doesn't know; his ID
# is 805347 (debut 2026-05-06), found by the eligibility pull 2026-10-12.
MLBAM_FIXES = {80537: 805347}

# Tab nicknames that don't end a team name in config/league.yaml. The workbook
# tab for the Absolute Sickos is still "Rats #8" (new regime in 2026).
TEAM_ALIASES: dict[str, str] = {"Rats": "Absolute Sickos"}

_DIVISION = re.compile(r"(AMERICAN|NATIONAL) LEAGUE\s+(.+?) DIVISION", re.S)


@dataclass
class WorkbookRosters:
    players: list[NplPlayer]
    tabs: dict[str, str]  # tab name -> config team name
    # Division the tab's header names, e.g. "AL Sultan of Swat", by team.
    divisions: dict[str, str] = field(default_factory=dict)

    def division_mismatches(self, alignment: Alignment) -> list[str]:
        """Teams whose tab header names a different division than the config."""
        config = {
            t: f"{lg.name} {div.name}"
            for lg in alignment.leagues
            for div in lg.divisions
            for t in div.teams
        }
        return [
            f"{team}: workbook {div}, config {config[team]}"
            for team, div in sorted(self.divisions.items())
            if team in config and div.upper() != config[team].upper()
        ]


def team_for_tab(tab: str, teams: list[str]) -> str | None:
    """Config team name for a tab like "Ducksnorts #12" or "Guys #21"."""
    nick = tab.split("#")[0].strip()
    if nick in TEAM_ALIASES:
        return TEAM_ALIASES[nick]
    matches = [t for t in teams if t.lower().endswith(nick.lower())]
    return matches[0] if len(matches) == 1 else None


def _name(raw: str) -> str:
    last, _, first = raw.partition(",")
    return f"{first.strip()} {last.strip()}".strip() if first else raw.strip()


def _flag(header: str) -> str | None | bool:
    """The roster flag a header starts, None for an active group, False if
    the row isn't a header."""
    h = header.strip().upper()
    if any(h.startswith(a) for a in ACTIVE_HEADERS):
        return None
    for prefix, flag in SECTIONS:
        if h.startswith(prefix):
            return flag
    return False


def read_tab(rows, team: str) -> tuple[list[NplPlayer], str | None]:
    """Players on one team tab, and the division its header names."""
    players: list[NplPlayer] = []
    division = None
    section: str | None | bool = False  # nothing until the first header
    for row in rows:
        cells = list(row) + [None] * 6
        a, b, c, d, pos, org = cells[:6]
        if division is None:
            for cell in cells:
                m = _DIVISION.search(str(cell or "").upper())
                if m:
                    lg = "AL" if m.group(1) == "AMERICAN" else "NL"
                    division = f"{lg} {m.group(2).strip().title().replace(' Of ', ' of ')}"
                    break
        if not isinstance(b, str) or not b.strip():
            continue
        if b.strip().upper().startswith(END_HEADER):
            break
        if not pos:
            flag = _flag(b)
            if flag is not False:
                section = flag
            continue
        if section is False or "," not in b:
            continue
        on_40 = isinstance(a, int | float) and a > 0
        flags = {f for f in ("roster_40man" if on_40 else None, section) if f}
        mlb_id = parse_mlb_id(int(d) if isinstance(d, float) else d)
        mlb_id = MLBAM_FIXES.get(mlb_id, mlb_id)
        ss = int(c) if isinstance(c, float) else c
        players.append(
            NplPlayer(
                mlb_id=mlb_id,
                name=_name(b),
                position=str(pos).strip(),
                mlb_org=str(org).strip() if org else None,
                team=team,
                flags=frozenset(flags),
                scoresheet_id=str(ss) if isinstance(ss, int | str) and str(ss).isdigit() else None,
            )
        )
    return players, division


def players_from_workbook(path: Path | str, alignment: Alignment) -> WorkbookRosters:
    """Every team tab's players. Raises if a tab can't be matched to a team
    or a config team has no tab."""
    from openpyxl import load_workbook  # only this interim source needs it

    teams = [t for lg in alignment.leagues for div in lg.divisions for t in div.teams]
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        out = WorkbookRosters(players=[], tabs={})
        unmatched = []
        for ws in wb.worksheets:
            if "#" not in ws.title:
                continue  # Key, 2026 Contracts, IDs, Service Time
            team = team_for_tab(ws.title, teams)
            if team is None:
                unmatched.append(ws.title)
                continue
            players, division = read_tab(ws.iter_rows(values_only=True), team)
            out.players += players
            out.tabs[ws.title] = team
            if division:
                out.divisions[team] = division
    finally:
        wb.close()
    missing = sorted(set(teams) - set(out.tabs.values()))
    if unmatched or missing:
        raise ValueError(f"unmatched tabs {unmatched}; teams without a tab {missing}")
    return out


def to_dump(players: list[NplPlayer]) -> list[dict]:
    """The players as a Django dumpdata export of the NPL app, which
    `sources.players_from_dump` reads (players without an MLBAM ID are left
    out: the dump is keyed by it)."""
    from npl_sim.rosters.npl_app import ROSTER_FLAGS

    teams = sorted({p.team for p in players})
    team_pk = {t: i + 1 for i, t in enumerate(teams)}
    records: list[dict] = [
        {"model": "npl.team", "pk": team_pk[t], "fields": {"full_name": t}} for t in teams
    ]
    for p in players:
        if p.mlb_id is None:
            continue
        records.append(
            {
                "model": "npl.player",
                "pk": str(p.mlb_id),
                "fields": {
                    "team": team_pk[p.team],
                    "name": p.name,
                    "position": p.position,
                    "mlb_org": p.mlb_org,
                    "scoresheet_id": p.scoresheet_id,
                    **{f: f in p.flags for f in ROSTER_FLAGS},
                },
            }
        )
    return records


def write_dump(players: list[NplPlayer], path: Path | str) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    records = to_dump(players)
    # One record per line keeps diffs readable when the rosters change.
    path.write_text("[\n" + ",\n".join(json.dumps(r) for r in records) + "\n]\n")
    return path
